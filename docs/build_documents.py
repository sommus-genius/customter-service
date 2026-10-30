from pathlib import Path
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
root=Path('docs'); source=(root/'项目审查与AI-Agent面试指南.md').read_text(encoding='utf-8-sig')
sections=re.split(r'(?m)^## ',source)
parts={x.splitlines()[0].split('.')[0]:x for x in sections[1:]}
intro='审查日期：2026年10月30日。本资料用于学习、改进项目与准备面试。已完成源码检查与部分本地故障探针，未完成真实模型、数据库与电商 API 的端到端验证。'
extra='''
## 先直接回答你的问题
以你描述的“前端开发，刚刚学完这个项目”的状态，我认为现在可以开始了解岗位和试投，但尚不足以把这个项目作为独立胜任 Agent 开发的证明。关键差距不是项目数量，而是你能否独立修改后端、处理失败、接通真实业务并验证效果。
再加一个 RAG 项目是合理的。两个项目分别展示“能办事”和“能基于资料回答问题”，但两个教程复现项目仍然不够。建议先把客服核心路径修通，再做带来源、拒答、更新和评测的 RAG，最后把两者集成。
你的目标可以优先选 AI 应用开发、Agent 应用开发和 AI 全栈岗位。前端经验可以用于流式交互、订单卡片、引用定位、工具执行进度和完整产品交付；同时要补齐 Python 异步、数据库事务、认证授权、工具调用和评测。当前材料不足以判断你能拿到什么职级或薪资。
## 第二个项目的具体交付标准
建议题目为“电商售后政策知识助手”。用户询问不同品类的退换货、运费及质保政策，系统返回有原文引用的答案；资料没有答案时明确拒答。订单是否发货仍走实时 API。
第一版先完成文档上传、结构化解析、按标题与段落分块、保存来源页码、向量召回和引用回答。前端支持查看原文与检索片段。不要在第一版同时引入多个数据库和复杂 Agent 编排。
第二版加入关键词与语义混合召回、重排、版本与权限过滤、文档删除更新。每新增一个模块，都用固定问题集比较正确率、召回率、延迟和成本，保留没有改善的案例。
作品交付至少包含：可运行代码、脱敏配置、启动说明、架构说明、演示视频、评测问题集与结果、典型失败分析。评测数据和业务改进数字必须来自实际运行。
## 什么时候开始投递
不用等到所有功能完善才投。客服核心流程稳定、能独立讲解和改功能后，可以开始试投并记录追问；RAG 做出可评测版本后扩大投递。若你仍依赖教程才能加一个 Action 或解释数据库并发，就先集中补这一块。
'''
def clean(s):
 s=re.sub(r'\[([^\]]+)\]\(([^)]+)\)',r'\1（\2）',s)
 return s.replace('**','').replace('`','')
def build(filename,title,text):
 d=Document(); sec=d.sections[0]; sec.top_margin=sec.bottom_margin=Cm(2); sec.left_margin=sec.right_margin=Cm(2.2)
 for name in ['Normal','Title','Heading 1','Heading 2','Heading 3']:
  st=d.styles[name]; st.font.name='Microsoft YaHei'; st._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); st.font.color.rgb=RGBColor(0,0,0)
 d.styles['Normal'].font.size=Pt(10.5); d.styles['Normal'].paragraph_format.space_after=Pt(6)
 d.styles['Normal'].paragraph_format.line_spacing=1.15
 d.add_paragraph(title,'Title'); d.add_paragraph(intro)
 for line in text.splitlines():
  line=line.strip()
  if not line: continue
  if line.startswith('|'):
   cells=[clean(x.strip()) for x in line.strip('|').split('|')]
   if all(re.fullmatch(r'[-: ]+',x) for x in cells): continue
   if cells[0] in ('优先级','阶段'): continue
   d.add_paragraph('；'.join(cells))
  elif line.startswith('#'):
   depth=len(line)-len(line.lstrip('#')); heading=clean(line.lstrip('# ').strip()); heading=re.sub(r'[^\w\s\u4e00-\u9fff]',' ',heading)
   d.add_paragraph(heading,'Heading 1' if depth<=2 else 'Heading 2')
  elif re.match(r'^\d+\. ',line) and '？' in line:
   q,a=(line.split('？',1)+[''])[:2]; d.add_paragraph(clean(q+'？'),'Heading 2'); d.add_paragraph(clean(a))
  else: d.add_paragraph(clean(line.lstrip('- ')))
 d.save(root/filename)
build('01 项目检查与改进报告.docx','电商客服项目检查与改进报告','## 检查结论\n'+parts['1']+'\n## 问题与改进\n'+parts['2'])
build('02 转岗评估与RAG项目准备.docx','前端转 AI Agent 岗位评估与准备',extra+'\n## 能力差距与准备路线\n'+parts['4']+'\n## 面试前自查\n'+parts['7'])
build('03 电商客服项目技术梳理.docx','电商客服项目架构与业务流程梳理','## 技术梳理\n'+parts['3']+'\n## 项目讲述\n'+parts['5'])
build('04 面试问题与参考回答.docx','电商客服与 RAG 面试问题及参考回答','本资料用于理解后练习表达。回答中的改进方案不代表当前代码已经实现；面试时应区分当前实现、发现的问题和个人改造。\n## 面试问题\n'+parts['6'])
print('Created four Word documents')
